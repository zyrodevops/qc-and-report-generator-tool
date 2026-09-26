"""
Unticked sections, rows and charts are left out of the report — in both the
HTML preview and the Word file — and nothing is deleted from the report data.
"""

import io

import docx

from app.render.docx.engine import render_docx
from app.render.html.engine import render_html
from app.seeds.defaults import get_default_block_state


def _state(commodity="APPLE"):
    st = get_default_block_state("perishable_sea_survey", commodity_key=commodity)
    st.setdefault("metadata", {})["number"] = "M-TEST-2026"
    return st


def _docx_text(state):
    d = docx.Document(io.BytesIO(render_docx(state)))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            parts.extend(c.text for c in r.cells)
    return "\n".join(parts), len(d.inline_shapes)


def test_every_text_section_starts_ticked():
    st = _state()
    narratives = [b for b in st["blocks"] if b["type"] == "narrative"]
    assert len(narratives) == 7
    assert all(b.get("included", True) is not False for b in narratives)


def test_unticked_section_is_left_out_of_html_and_docx():
    st = _state()
    note = next(b for b in st["blocks"] if b["id"] == "b_note")
    heading = note["section"]
    assert heading in render_html(st)

    note["included"] = False
    assert heading not in render_html(st)
    text, _ = _docx_text(st)
    assert heading not in text
    # still there in the data, ready to be ticked back in
    assert any(b["id"] == "b_note" for b in st["blocks"])


def test_measurement_rows_follow_their_ticks():
    st = _state("BLUEBERRY")  # berries: no penetrometer by default
    m = next(b for b in st["blocks"] if b["type"] == "measurements")
    subjects = {r["subject"]: r["included"] for r in m["rows"]}
    assert subjects["Pulp Temperature"] is True
    assert subjects["Fruit Pressure"] is False
    html = render_html(st)
    assert "Pulp Temperature" in html and "Fruit Pressure" not in html

    # the surveyor ticks it in for a rare report
    next(r for r in m["rows"] if r["subject"] == "Fruit Pressure")["included"] = True
    assert "Fruit Pressure" in render_html(st)


def test_only_real_measurements_are_offered():
    m = next(b for b in _state()["blocks"] if b["type"] == "measurements")
    assert [r["subject"] for r in m["rows"]] == ["Pulp Temperature", "Brix", "Fruit Pressure"]


def test_table_heading_and_chart_follow_their_ticks():
    st = _state("APPLE")
    t = next(b for b in st["blocks"] if b["type"] == "table")
    t["rows"] = [{"group": "120", "values": {"sound": "67", "rotten": "3"}}]

    t["show_title"], t["show_chart"] = True, True
    html = render_html(st)
    assert t["title"] in html and "SURVEY FINDINGS IN GRAPH" in html

    t["show_title"], t["show_chart"] = False, False
    html = render_html(st)
    assert t["title"] not in html and "SURVEY FINDINGS IN GRAPH" not in html
    # the table itself always stays
    assert "Total (pcs)" in html


def test_old_reports_without_flags_render_as_before():
    st = _state()
    for b in st["blocks"]:
        b.pop("included", None)
        b.pop("show_title", None)
        b.pop("show_chart", None)
        for r in b.get("rows", []) or []:
            if isinstance(r, dict):
                r.pop("included", None)
    html = render_html(st)
    for b in st["blocks"]:
        if b["type"] == "narrative":
            assert b["section"] in html
