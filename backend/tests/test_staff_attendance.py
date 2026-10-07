"""
Attendance from the private staff list, and the reference-data endpoint.

The real lists are not in git (backend/private, see app.seeds.private_data),
so these tests write made-up lists to a temporary folder.
"""

import json
import os

import docx
import pytest
from fastapi.testclient import TestClient

from app.config import settings
from app.main import app
from app.render.docx.engine import render_narrative
from app.render.html.engine import render_narrative_html
from app.seeds import private_data, staff_lookup
from app.seeds.staff_lookup import get_default_attendance, lookup_consignee_staff

FIRM = "Marine Cargo Agencies Pvt.Ltd (On behalf of Consignees)"

STAFF = {
    "consignees": [
        {"company": "Alpha Orchards Pvt. Ltd", "name": "Mr. Test One", "designation": "Sales Manager",
         "representing": "Alpha Orchards Pvt. Ltd - (Consignees)"},
        {"company": "Beta Retail Limited", "name": "Mr. Test Two", "designation": "QC Manager",
         "representing": "Beta Retail Limited - (Consignees)"},
        {"company": "Beta Retail Limited", "name": "Mr. Test Three", "designation": "Store Manager",
         "representing": "Beta Retail Limited - (Consignees)"},
        {"company": "Gamma & Co", "name": "Mr. Test Four", "designation": "Owner",
         "representing": "Gamma & Co - (Consignees)"},
    ],
    "shipping_lines": [],
    "shippers": [],
    "cargo_insurers": [],
    "mca_surveyors": [
        {"name": "Mr. Firm Surveyor", "designation": "Surveyor", "representing": FIRM},
        {"name": "Mr. Second Surveyor", "designation": "Surveyor", "representing": FIRM},
    ],
}
COLD_STORAGES = [
    {"id": "cs_001", "city": "TESTCITY", "name": "Test Cold Store (Unit 1)",
     "clean_name": "Test Cold Store", "address": "1 Test Road, Test City 400001"},
]


@pytest.fixture
def private_lists(tmp_path, monkeypatch):
    (tmp_path / private_data.STAFF_FILE).write_text(json.dumps(STAFF), encoding="utf-8")
    (tmp_path / private_data.COLD_STORAGE_FILE).write_text(json.dumps(COLD_STORAGES), encoding="utf-8")
    monkeypatch.setattr(settings, "PRIVATE_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(staff_lookup, "_CACHE", None)
    yield tmp_path
    staff_lookup._CACHE = None


def test_lookup_consignee_staff(private_lists):
    alpha = lookup_consignee_staff("Alpha Orchards")
    assert [m["name"] for m in alpha] == ["Mr. Test One"]

    beta = lookup_consignee_staff("M/s Beta Retail Limited, Navi Mumbai")
    assert {m["name"] for m in beta} == {"Mr. Test Two", "Mr. Test Three"}

    gamma = lookup_consignee_staff("Gamma & Co.")
    assert [m["name"] for m in gamma] == ["Mr. Test Four"]

    assert lookup_consignee_staff("Unknown Importers") == []
    assert lookup_consignee_staff("[Consignee]") == []


def test_get_default_attendance(private_lists):
    # Without consignee: the firm's first surveyor only
    att_default = get_default_attendance("")
    assert att_default == [{"name": "Mr. Firm Surveyor", "designation": "Surveyor", "representing": FIRM}]

    # With consignee: their staff, then the firm's surveyor
    att = get_default_attendance("Alpha Orchards Pvt. Ltd")
    assert [r["name"] for r in att] == ["Mr. Test One", "Mr. Firm Surveyor"]


def test_missing_lists_give_empty_attendance(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "PRIVATE_DATA_DIR", str(tmp_path / "absent"))
    monkeypatch.setattr(staff_lookup, "_CACHE", None)
    assert get_default_attendance("Alpha Orchards") == []
    assert private_data.load_cold_storages() == []
    staff_lookup._CACHE = None


def test_reference_data_needs_login(private_lists):
    client = TestClient(app)
    assert client.get("/api/reference-data").status_code == 401

    login = client.post("/api/auth/login", json={"email": "surveyor@example.com", "password": "Password123!"})
    assert login.status_code == 200
    res = client.get("/api/reference-data", headers={"Authorization": f"Bearer {login.json()['token']}"})
    assert res.status_code == 200
    data = res.json()
    assert data["staff"]["mca_surveyors"][0]["name"] == "Mr. Firm Surveyor"
    assert data["cold_storages"][0]["clean_name"] == "Test Cold Store"


def test_banner_is_rechecked_not_cached_stale(tmp_path, monkeypatch):
    from PIL import Image

    monkeypatch.setattr(settings, "PRIVATE_DATA_DIR", str(tmp_path))
    Image.new("RGBA", (400, 80), (0, 32, 96, 255)).save(tmp_path / private_data.BANNER_FILE)
    client = TestClient(app)
    login = client.post("/api/auth/login", json={"email": "surveyor@example.com", "password": "Password123!"})
    auth = {"Authorization": f"Bearer {login.json()['token']}"}

    res = client.get("/api/reference-data/banner.jpg", headers=auth)
    assert res.status_code == 200 and res.headers["content-type"] == "image/jpeg"
    assert "no-cache" in res.headers["cache-control"]
    etag = res.headers["etag"]
    # unchanged: the browser keeps its copy
    assert client.get("/api/reference-data/banner.jpg", headers={**auth, "If-None-Match": etag}).status_code == 304
    # a new banner file: the browser gets it at once
    Image.new("RGBA", (400, 80), (200, 0, 0, 255)).save(tmp_path / private_data.BANNER_FILE)
    os.utime(tmp_path / private_data.BANNER_FILE, (1, 1))
    res = client.get("/api/reference-data/banner.jpg", headers={**auth, "If-None-Match": etag})
    assert res.status_code == 200 and res.headers["etag"] != etag


def test_render_narrative_docx_and_html():
    block = {
        "id": "b_para1",
        "type": "narrative",
        "section": "PARAGRAPH 1: APPLICATION",
        "additional_text": "Pursuant to the Consignee's request and subsequent appointment...",
        "attendance_intro": "The following persons attended the survey:",
        "attendance": [
            {"name": "Mr. Test One", "designation": "Sales Manager",
             "representing": "Alpha Orchards Pvt. Ltd - (Consignees)"},
            {"name": "Mr. Firm Surveyor", "designation": "Surveyor", "representing": FIRM},
        ],
    }

    # DOCX
    doc = docx.Document()
    render_narrative(doc, block)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    headers = [c.text for c in table.rows[0].cells]
    assert headers == ["Name", "Designation", "Representing"]
    row1 = [c.text for c in table.rows[1].cells]
    assert "Test One" in row1[0]
    assert "Sales Manager" in row1[1]

    # HTML
    html = render_narrative_html(block)
    assert "PARAGRAPH 1: APPLICATION" in html
    assert "The following persons attended the survey:" in html
    assert "preview-table" in html
    assert "Test One" in html
    assert "Firm Surveyor" in html
