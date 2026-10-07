"""
The client's report look on survey reports (app/render/house_style.py, app/render/docx/house.py):
A4 with a page border, the title "FINAL SURVEY REPORT NO. …", headings bold / underlined /
dark blue, **bold** and __underline__ marks, the findings table's "Percentage" row, and the
closing (disclaimer, issued without prejudice, licence line, signature spaces, ØØØ).
"""

import io

import docx
from docx.oxml.ns import qn

from app.config import settings
from app.render import house_style as hs
from app.render.docx.engine import render_docx
from app.render.narrative_text import split_narrative
from app.seeds.defaults import get_default_block_state


def _apple(number="M-321-2026"):
    st = get_default_block_state("perishable_sea_survey", commodity_key="APPLE")
    st["metadata"]["number"] = number
    return st


def _doc(st):
    return docx.Document(io.BytesIO(render_docx(st)))


def test_title_heading_and_marks():
    assert hs.report_title({"metadata": {"number": "M-109-2026"}}) == "FINAL SURVEY REPORT NO. M-109-2026"
    assert hs.report_title({"metadata": {"number": "M-5-2026", "state": "PRELIMINARY"}}) == "PRELIMINARY SURVEY REPORT NO. M-5-2026"
    assert hs.heading_text("PARAGRAPH 1: APPLICATION") == "PARAGRAPH 1: APPLICATION:"
    assert hs.heading_text("NOTE:") == "Note:"
    assert hs.runs("a **b __c__** d") == [("a ", False, False), ("b ", True, False), ("c", True, True), (" d", False, False)]
    assert hs.plain("## Head:\n**x** __y__") == "Head:\nx y"
    parts = split_narrative("## Sub:\nText.\n\n1) FRL **(70 Count):** x\n2) VBT **(80 Count):** y\n• a\n• b", rich=True)
    assert parts == [("h", ["Sub:"]), ("p", ["Text."]),
                     ("ol", ["1) FRL **(70 Count):** x", "2) VBT **(80 Count):** y"]), ("ul", ["a", "b"])]
    # the old callers see what they always did
    assert split_narrative("1) one\n2) two") == [("p", ["1) one", "2) two"])]


def test_banner_is_in_the_body_not_the_header(tmp_path, monkeypatch):
    # Word greys out header pictures while the body is edited, so the banner
    # is pinned to page 1 in the body, above the title, with the text below it.
    from PIL import Image

    monkeypatch.setattr(settings, "PRIVATE_DATA_DIR", str(tmp_path))
    Image.new("RGB", (400, 80), (0, 32, 96)).save(tmp_path / "banner.png")
    d = _doc(_apple())
    s = d.sections[0]
    assert not s.first_page_header._element.xpath(".//w:drawing")
    title = d.paragraphs[0]
    assert title.text == "FINAL SURVEY REPORT NO. M-321-2026"
    anchor = title._p.find(".//" + qn("wp:anchor"))
    assert anchor is not None and anchor.find(qn("wp:wrapTopAndBottom")) is not None
    for tag in ("wp:positionH", "wp:positionV"):
        pos = anchor.find(qn(tag))
        assert pos.get("relativeFrom") == "page"
        assert abs(int(pos.find(qn("wp:posOffset")).text) / 360000 - hs.BORDER_CM) < 0.01


def test_closing_parts():
    assert hs.dated_parts({"dated": "2026-09-23"}) == ("23", "rd", " September 2026.")
    assert hs.dated_parts({"content": "“ISSUED”\nDated: 11 October 2026\n"}) == ("11", "th", " October 2026.")
    assert hs.licence_line("") == "(IRDA Surveyor License No.): [License No.]."
    assert hs.licence_line("ABC/1.") == "(IRDA Surveyor License No.): ABC/1."
    st = {"blocks": [{"attendance": [{"name": "Mr. A. Kumar & Team", "representing": "Marine Cargo Agencies Pvt.Ltd"}]}]}
    assert hs.reporting_surveyor(st) == "A. Kumar"


def test_word_page_title_order_and_closing(monkeypatch):
    monkeypatch.setattr(settings, "IRDAI_LICENCE_NUMBER", "")
    st = _apple()
    for b in st["blocks"]:
        if b["id"] == "b_closure":
            b["dated"] = "2026-09-28"
    d = _doc(st)
    s = d.sections[0]
    assert round(s.page_width.cm, 1) == 21.0 and round(s.page_height.cm, 1) == 29.7
    assert s._sectPr.find(qn("w:pgBorders")) is not None
    assert s.different_first_page_header_footer
    assert "FINAL SURVEY REPORT NO. M-321-2026" in s.header.paragraphs[0].text

    paras = [p.text for p in d.paragraphs]
    assert paras[0] == "FINAL SURVEY REPORT NO. M-321-2026"
    title_run = d.paragraphs[0].runs[0]
    assert title_run.bold and title_run.underline and title_run.font.size.pt == 14
    assert "PARAGRAPH 1: APPLICATION:" in paras
    assert "Note:" not in paras  # an empty note prints nothing
    assert "THE CONDITION FOUND OF APPLE FRUITS:" in paras
    assert not any("**" in p or "__" in p or p.startswith("## ") for p in paras)
    # the closing comes after Paragraph 5, in the client's order and wording
    i = paras.index(hs.DISCLAIMER_HEADING)
    assert i > paras.index("PARAGRAPH 5: DOCUMENTATION:")
    assert paras[i + 1] == hs.DISCLAIMER
    assert paras[i + 2:i + 6] == [hs.ISSUED, "Dated: 28th September 2026.", "Place: Mumbai, India.",
                                  "(IRDA Surveyor License No.): [License No.]."]
    assert paras[-1] == hs.END_MARK
    sig = d.tables[-1]
    assert hs.SIGNATURE_LABEL in sig.cell(0, 1).text and hs.COMPANY in sig.cell(0, 1).text


def test_bold_and_underline_runs_in_word():
    d = _doc(_apple())
    runs = {r.text: r for p in d.paragraphs for r in p.runs}
    assert runs["Sound apples:"].bold
    assert runs["The Apple fruits were cut, and the following pulp conditions were observed"].underline
    assert not runs[":"].underline if ":" in runs else True


def test_findings_table_percentage_row_and_shading():
    st = _apple()
    t = next(b for b in st["blocks"] if b["id"] == "b_table")
    cats = [c["key"] for c in t["categories"]][:3]
    t["rows"] = [{"group": g, "values": {k: "1" for k in cats}} for g in ("90", "100", "110")]
    d = _doc(st)
    table = next(tb for tb in d.tables if tb.rows[-1].cells[0].text == "Percentage")
    def fill(cell):
        shd = cell._tc.tcPr.find(qn("w:shd")) if cell._tc.tcPr is not None else None
        return shd.get(qn("w:fill")) if shd is not None else None
    assert fill(table.rows[0].cells[0]) == hs.GREY_HEAD     # header
    assert fill(table.rows[2].cells[0]) == hs.GREY_ALT      # second data row
    assert fill(table.rows[1].cells[0]) is None             # first data row
    assert fill(table.rows[-1].cells[0]) == hs.GREY_HEAD    # percentage row


def test_client_marks_are_applied_to_plain_text():
    """Text written before the marks existed (or typed) gets the client's bold / underline when printed."""
    apple = (
        "THE CONDITION FOUND OF APPLE FRUITS:\n\n"
        "• Royal Gala (135 Count): Range of 14.18 LBS to 17.59 LBS.\n\n"
        "The Apple fruits were cut, and the following pulp conditions were observed:\n"
        "• Sound apples: The pulp was consistently hard and white.\n"
        "• The sugar brix was measured.\n\n"
        "During our inspection, 16 boxes out of the 8,232 boxes (under 5 counts) were separated."
    )
    assert hs.auto_marks(apple, "APPLE").split("\n") == [
        "## THE CONDITION FOUND OF APPLE FRUITS:", "",
        "• **Royal Gala (135 Count):** Range of 14.18 LBS to 17.59 LBS.", "",
        "__The Apple fruits were cut, and the following pulp conditions were observed__:",
        "• **Sound apples:** The pulp was consistently hard and white.",
        "• The sugar brix was measured.", "",
        "During our inspection, **16 boxes out of the 8,232 boxes** (under 5 counts) were separated.",
    ]
    # Pear prints only the count bold (M-162)
    assert hs.auto_marks("1) FRL (70 Count): Range of 11.91 LBS.", "pear") == "1) FRL **(70 Count):** Range of 11.91 LBS."
    cause = (
        "Additional contributing factors observed during the inspection include:\n"
        "• Mechanical injury, likely sustained during sorting.\n"
        "• Pressure damage (bruising) indicative of improper handling."
    )
    assert hs.auto_marks(cause, "apple").split("\n") == [
        "__Additional contributing factors observed during the inspection include__:",
        "• **Mechanical injury**, likely sustained during sorting.",
        "• **Pressure damage (bruising)** indicative of improper handling.",
    ]
    assert hs.auto_marks("Findings & Assessment:", "mandarin") == "## Findings & Assessment:"
    assert hs.auto_marks("• Temperature Data Recorder - Annexure A1 & A2", "mandarin") == "• Temperature Data Recorder - **Annexure A1 & A2**"
    # already marked, another fruit, QC: left alone
    assert hs.auto_marks("• **Sound apples:** x", "apple") == "• **Sound apples:** x"
    assert hs.auto_marks("• Sound apples: x", "kiwi") == "• Sound apples: x"


def test_old_plain_text_prints_with_marks_in_word():
    st = _apple()
    p21 = next(b for b in st["blocks"] if b["id"] == "b_para2_1")
    p21["additional_text"] = hs.plain(p21["additional_text"])  # an older report: no marks in its text
    d = _doc(st)
    runs = {r.text: r for p in d.paragraphs for r in p.runs}
    assert runs["Sound apples:"].bold
    assert runs["THE CONDITION FOUND OF APPLE FRUITS:"].underline


def test_qc_reports_keep_their_own_look():
    st = get_default_block_state("perishable_qc_sea", commodity_key="APPLE")
    st["metadata"]["number"] = "M-322-2026"
    d = _doc(st)
    paras = [p.text for p in d.paragraphs]
    assert "FINAL SURVEY REPORT NO. M-322-2026" not in paras
    assert hs.DISCLAIMER_HEADING not in paras
