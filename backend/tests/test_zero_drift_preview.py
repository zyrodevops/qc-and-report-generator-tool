"""
Automated Zero-Drift Verification Test Suite for Milestone 1.
Verifies that Word .docx generation and HTML preview produce bit-for-bit
identical calculations, percentages, row totals, column totals, and photo ranges.
"""

from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Set
from bs4 import BeautifulSoup
import docx
import pytest

from app.render.docx.engine import render_docx
from app.render.html.engine import render_html
from tests.e2e.helpers.synthetic_data import make_synthetic_block_state


def extract_docx_data(docx_bytes: bytes) -> Dict[str, Any]:
    """Extracts tables, defect metrics, captions, and range strings from DOCX."""
    doc = docx.Document(io.BytesIO(docx_bytes))

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    tables = []
    for t in doc.tables:
        t_rows = []
        for r in t.rows:
            t_rows.append([c.text.strip() for c in r.cells])
        tables.append(t_rows)

    defect_tables = []
    for t_rows in tables:
        if not t_rows:
            continue
        header = t_rows[0]
        if any("Total (" in h for h in header):
            data_rows = t_rows[1:-2]
            totals_row = t_rows[-2]
            pct_row = t_rows[-1]

            defect_tables.append({
                "header": header,
                "data_rows": data_rows,
                # Last column is the row total. The old final "%" column,
                # which joined every row percentage with " / " and stretched
                # each row several times its height, was removed; per-column
                # percentages are in the foot row, suffixed with "%".
                "row_totals": [r[-1] for r in data_rows],
                "row_pcts": [],
                "col_totals": totals_row[1:-1],
                "grand_total": totals_row[-1],
                "col_pcts": [c.rstrip("%") for c in pct_row[1:-1]],
                "full_grid": t_rows,
            })

    photo_captions = []
    for p in doc.paragraphs:
        m = re.match(r"^Survey Photo No\. \d+(?: — .*)?$",p.text.strip())
        if m:
            photo_captions.append(m.group(0))
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    m = re.match(r"^Survey Photo No\. \d+(?: — .*)?$",p.text.strip())
                    if m:
                        photo_captions.append(m.group(0))

    combined_text = " ".join(paragraphs)
    for t_rows in tables:
        for r in t_rows:
            combined_text += " " + " ".join(r)

    photo_ranges = re.findall(r"\(Photo Nos?\. [^\)]+\)", combined_text)
    numeric_tokens = set(re.findall(r"\b\d+(?:\.\d+)?%?\b", combined_text))

    return {
        "paragraphs": paragraphs,
        "tables": tables,
        "defect_tables": defect_tables,
        "photo_captions": photo_captions,
        "photo_ranges": photo_ranges,
        "numeric_tokens": numeric_tokens,
    }


def extract_html_data(html_content: str) -> Dict[str, Any]:
    """Extracts tables, defect metrics, captions, and range strings from HTML preview."""
    soup = BeautifulSoup(html_content, "html.parser")

    paragraphs = [
        p.get_text(strip=True)
        for p in soup.find_all(["p", "h1", "h2", "h3", "h4", "h5", "h6"])
        if p.get_text(strip=True)
    ]

    tables = []
    for t in soup.find_all("table"):
        t_rows = []
        for tr in t.find_all("tr"):
            cells = [cell.get_text(strip=True) for cell in tr.find_all(["th", "td"])]
            if cells:
                t_rows.append(cells)
        if t_rows:
            tables.append(t_rows)

    defect_tables = []
    for t_rows in tables:
        if not t_rows:
            continue
        header = t_rows[0]
        if any("Total (" in h for h in header):
            data_rows = t_rows[1:-2]
            totals_row = t_rows[-2]
            pct_row = t_rows[-1]

            defect_tables.append({
                "header": header,
                "data_rows": data_rows,
                # Last column is the row total. The old final "%" column,
                # which joined every row percentage with " / " and stretched
                # each row several times its height, was removed; per-column
                # percentages are in the foot row, suffixed with "%".
                "row_totals": [r[-1] for r in data_rows],
                "row_pcts": [],
                "col_totals": totals_row[1:-1],
                "grand_total": totals_row[-1],
                "col_pcts": [c.rstrip("%") for c in pct_row[1:-1]],
                "full_grid": t_rows,
            })

    photo_captions = []
    for text in soup.stripped_strings:
        m = re.match(r"^Survey Photo No\. \d+(?: — .*)?$",text)
        if m:
            photo_captions.append(m.group(0))

    full_text = soup.get_text(separator=" ", strip=True)
    photo_ranges = re.findall(r"\(Photo Nos?\. [^\)]+\)", full_text)
    numeric_tokens = set(re.findall(r"\b\d+(?:\.\d+)?%?\b", full_text))

    return {
        "paragraphs": paragraphs,
        "tables": tables,
        "defect_tables": defect_tables,
        "photo_captions": photo_captions,
        "photo_ranges": photo_ranges,
        "numeric_tokens": numeric_tokens,
    }


def assert_zero_drift(docx_bytes: bytes, html_content: str) -> None:
    """Rigorous differential validator between DOCX and HTML."""
    docx_data = extract_docx_data(docx_bytes)
    html_data = extract_html_data(html_content)

    # 1. Defect table count equality
    assert len(docx_data["defect_tables"]) == len(html_data["defect_tables"]), (
        f"Defect table count mismatch: DOCX has {len(docx_data['defect_tables'])}, "
        f"HTML has {len(html_data['defect_tables'])}"
    )

    # 2. Defect table cell-for-cell equality
    for idx, (d_tab, h_tab) in enumerate(zip(docx_data["defect_tables"], html_data["defect_tables"])):
        assert d_tab["header"] == h_tab["header"], f"Table {idx} header drift"
        assert d_tab["row_totals"] == h_tab["row_totals"], f"Table {idx} row totals drift"
        assert d_tab["row_pcts"] == h_tab["row_pcts"], f"Table {idx} row percentages drift"
        assert d_tab["col_totals"] == h_tab["col_totals"], f"Table {idx} column totals drift"
        assert d_tab["grand_total"] == h_tab["grand_total"], f"Table {idx} grand total drift"
        assert d_tab["col_pcts"] == h_tab["col_pcts"], f"Table {idx} column percentages drift"
        assert d_tab["full_grid"] == h_tab["full_grid"], f"Table {idx} full grid drift"

    # 3. Photo captions equality
    assert docx_data["photo_captions"] == html_data["photo_captions"], "Photo captions drift"

    # 4. Photo range strings equality
    assert docx_data["photo_ranges"] == html_data["photo_ranges"], "Photo range strings drift"


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

def test_zero_drift_mandarin_citrus_benchmark():
    """Validates Mandarin QC 8-box benchmark (pcs, 2 decimal places, Hare-Niemeyer)."""
    state = make_synthetic_block_state(mode="SEA")

    docx_bytes = render_docx(state)
    html_content = render_html(state)

    assert_zero_drift(docx_bytes, html_content)

    # Specific benchmark values check
    html_data = extract_html_data(html_content)
    tab = html_data["defect_tables"][0]
    assert tab["row_totals"] == ["234"]
    # No joined-percentage column: every row is exactly as wide as the header.
    assert "%" not in tab["header"]
    assert all(len(r) == len(tab["header"]) for r in tab["data_rows"])
    assert tab["col_totals"] == ["133", "54", "14", "24", "9"]
    assert tab["grand_total"] == "234"
    assert tab["col_pcts"] == ["56.84", "23.08", "5.98", "10.26", "3.84"]


def test_zero_drift_table_grapes_kg_3dp_benchmark():
    """Validates Table Grapes benchmark (kg unit, 3 decimal places)."""
    grapes_state = {
        "metadata": {
            "number": "M-002-2026",
            "family": "QC_REPORT",
            "docx_template": "mca-qc-synthetic.docx",
        },
        "blocks": [
            {
                "id": "b_grapes",
                "type": "table",
                "title": "Table Grapes Defect Analysis",
                "unit": "kg",
                "grouping_label": "Sample Batch",
                "categories": [
                    {"key": "sound", "label": "Sound (kg)"},
                    {"key": "waterberry", "label": "Waterberry (kg)"},
                    {"key": "decay", "label": "Decay (kg)"},
                ],
                "rows": [
                    {"group": "Batch 1", "values": {"sound": "5.190", "waterberry": "0.606", "decay": "0.434"}},
                    {"group": "Batch 2", "values": {"sound": "42.372", "waterberry": "1.658", "decay": "0.442"}},
                ],
            }
        ],
    }

    docx_bytes = render_docx(grapes_state)
    html_content = render_html(grapes_state)

    assert_zero_drift(docx_bytes, html_content)

    html_data = extract_html_data(html_content)
    tab = html_data["defect_tables"][0]
    # Assert 3 decimal place formatting
    assert tab["row_totals"] == ["6.230", "44.472"]
    assert tab["col_totals"] == ["47.562", "2.264", "0.876"]
    assert tab["grand_total"] == "50.702"
    assert tab["col_pcts"] == ["93.81", "4.46", "1.73"]


def test_zero_drift_photo_ranges_multi_group():
    """Validates single, pair, and range bracketed photo references."""
    photo_state = {
        "metadata": {"number": "M-003-2026", "docx_template": "mca-qc-synthetic.docx"},
        "blocks": [
            {
                "id": "b_narr",
                "type": "narrative",
                "section": "EVIDENCE",
                "additional_text": "Initial inspection (Photo No. 1), container seals (Photo Nos. 2 & 3), and cartons (Photo Nos. 4 to 7).",
            },
            {
                "id": "b_plate",
                "type": "photo_plate",
                "label": "Photographic Plate",
                "groups": [
                    {"id": "g1", "observation": "Container exterior", "asset_ids": ["a1"]},
                    {"id": "g2", "observation": "Customs seal intact", "asset_ids": ["a2", "a3"]},
                    {"id": "g3", "observation": "Defective fruit cartons", "asset_ids": ["a4", "a5", "a6", "a7"]},
                ],
            },
        ],
    }

    docx_bytes = render_docx(photo_state)
    html_content = render_html(photo_state)

    assert_zero_drift(docx_bytes, html_content)

    docx_data = extract_docx_data(docx_bytes)
    html_data = extract_html_data(html_content)

    expected_ranges = ["(Photo No. 1)", "(Photo Nos. 2 & 3)", "(Photo Nos. 4 to 7)"]
    assert docx_data["photo_ranges"] == expected_ranges
    assert html_data["photo_ranges"] == expected_ranges
