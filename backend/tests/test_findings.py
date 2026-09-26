"""
The graph and the FINAL SUMMARY under a tally table, as the client draws them.

His reports use a column chart (144 of 180 graph pages), coloured by what each
column is, with a Total column at 100% (134 of 180) and the title under it;
and, when the counts come in groups, a FINAL SUMMARY of each group and the whole.
"""

import io

import docx

from app.compute.arithmetic import compute
from app.compute.traceability import check_traceability
from app.render import findings
from app.render.docx.engine import render_docx
from app.render.html.engine import render_html

C1, C2 = "TSTU1234565", "TSTU7654326"


def _table(**extra):
    return {
        "id": "t", "type": "table", "title": "CONDITION FOUND", "unit": "pcs", "grouping_label": "Count",
        "categories": [{"key": "sound", "label": "Sound (Pcs)"}, {"key": "pressed", "label": "Pressed Marks"},
                       {"key": "rot", "label": "Rotten"}, {"key": "creased", "label": "Creased"}],
        "rows": [
            {"group": "88", "container": C1, "boxes_opened": 1, "values": {"sound": "72", "pressed": "8", "rot": "8", "creased": "0"}},
            {"group": "88", "container": C1, "boxes_opened": 1, "values": {"sound": "57", "pressed": "23", "rot": "8", "creased": "0"}},
            {"group": "113", "container": C2, "boxes_opened": 1, "values": {"sound": "80", "pressed": "0", "rot": "8", "creased": "0"}},
        ],
        "show_chart": True,
        **extra,
    }


def _computed(block):
    st = compute({"blocks": [block]})
    return st["blocks"][0], st["blocks"][0]["_computed"]


def test_bars_as_the_client_draws_them():
    # The weights from one of the client's blueberry reports; his chart reads
    # 5.25%, 64.37%, 30.38% and a Total at 100%.
    kg = {"id": "k", "type": "table", "unit": "kg",
          "categories": [{"key": "sound", "label": "Sound"}, {"key": "soft", "label": "Soft"}, {"key": "rot", "label": "Rotten"}],
          "rows": [{"group": "1", "values": {"sound": "0.393", "soft": "4.815", "rot": "2.272"}}]}
    b, c = _computed(kg)
    assert findings.chart_bars(b, c) == [
        ("Sound (0.393 Kg)", 5.25, findings.GREEN),
        ("Soft (4.815 Kg)", 64.37, findings.YELLOW),
        ("Rotten (2.272 Kg)", 30.38, findings.RED),
        ("Total (7.480 Kg)", 100.0, findings.BLUE),
    ]
    assert findings.chart_title(b) == "LOSS CALCULATION IN GRAPH"


def test_counts_titles_and_empty_columns():
    b, c = _computed(_table())
    bars = findings.chart_bars(b, c)
    # Creased has no figure anywhere: no column for it. The "(Pcs)" of the
    # heading is the unit, not part of the name.
    assert [x[0] for x in bars] == ["Sound (209 Pcs)", "Pressed Marks (31 Pcs)", "Rotten (24 Pcs)", "Total (264 Pcs)"]
    assert findings.chart_title(b) == "SURVEY FINDINGS IN GRAPH"
    assert findings.chart_title({**b, "chart_title": "  Overall survey findings in graph "}) == "Overall survey findings in graph"
    assert [x[0] for x in findings.chart_bars({**b, "chart_total_bar": False}, c)][-1] == "Rotten (24 Pcs)"
    png = findings.chart_png(b, c)
    assert png and png.startswith(bytes([0x89]) + b"PNG")


def test_final_summary_by_container():
    b, c = _computed(_table(summary={"show": True, "by": "container"}))
    s = c["summary"]
    assert [(g["key"], g["boxes"], str(g["grand_total"])) for g in s["groups"]] == [(C1, 2, "176"), (C2, 1, "88")]
    assert str(s["groups"][0]["column_percentages"]["sound"]) == "73.30"
    rows = findings.summary_rows(b, c)
    assert rows["title"] == "FINAL SUMMARY"
    assert rows["header"][0] == "Container" and rows["header"][-1] == "Total (pcs)"
    assert [r["cells"][0] for r in rows["rows"]] == [
        f"{C1} (2 Boxes)", "Percentage", f"{C2} (1 Box)", "Percentage", "Total 3 Boxes", "Percentage"]
    assert rows["rows"][4]["cells"][-1] == str(c["grand_total"])

    # By count instead; and nothing when it is not ticked or has one group.
    b2, c2 = _computed(_table(summary={"show": True, "by": "group"}))
    assert [g["key"] for g in c2["summary"]["groups"]] == ["88", "113"]
    b3, c3 = _computed(_table())
    assert "summary" not in c3 and findings.summary_rows(b3, c3) is None
    one = _table(summary={"show": True, "by": "container"})
    one["rows"] = one["rows"][:2]
    assert findings.summary_rows(*_computed(one)) is None


def test_word_and_html_carry_the_summary_and_the_graph():
    state = {"metadata": {"family": "QC_REPORT"}, "blocks": [_table(summary={"show": True, "by": "container"})]}
    raw = render_docx(state)
    d = docx.Document(io.BytesIO(raw))
    texts = [p.text for p in d.paragraphs]
    assert "FINAL SUMMARY" in texts
    summary = d.tables[-1]
    assert summary.rows[0].cells[0].text == "Container"
    assert summary.rows[-2].cells[0].text == "Total 3 Boxes"
    assert len(d.inline_shapes) == 1  # the graph
    # Every figure in the summary traces back to the cells.
    assert check_traceability(state, raw).passed

    html = render_html(state)
    assert "FINAL SUMMARY" in html and f"{C2} (1 Box)" in html
    assert 'alt="SURVEY FINDINGS IN GRAPH"' in html

    state["blocks"][0]["show_chart"] = False
    assert len(docx.Document(io.BytesIO(render_docx(state))).inline_shapes) == 0


def test_counts_are_whole_weights_keep_three_places():
    b, c = _computed(_table())
    assert [str(v) for v in c["row_totals"]] == ["88", "88", "88"] and str(c["grand_total"]) == "264"
    kg = {"id": "k", "type": "table", "unit": "kg", "categories": [{"key": "a", "label": "Sound"}],
          "rows": [{"group": "1", "values": {"a": "0.82"}}]}
    assert str(_computed(kg)[1]["grand_total"]) == "0.820"
    # A fraction in a count table is not rounded out of sight.
    odd = _table()
    odd["rows"][0]["values"]["sound"] = "72.5"
    assert str(_computed(odd)[1]["row_totals"][0]) == "88.50"


def test_a_cleared_cell_does_not_stop_the_report():
    t = _table()
    t["rows"][0]["values"]["pressed"] = ""
    t["rows"][1]["values"]["rot"] = "  "
    b, c = _computed(t)
    assert str(c["row_totals"][0]) == "80" and str(c["row_totals"][1]) == "80"
    assert render_docx({"metadata": {}, "blocks": [t]})


def test_a_graph_with_many_columns_still_draws():
    cats = [{"key": f"c{i}", "label": f"Defect Name Number {i} (Pcs)"} for i in range(11)]
    t = {"id": "w", "type": "table", "unit": "pcs", "categories": cats, "show_chart": True,
         "rows": [{"group": "100", "values": {f"c{i}": str(i + 1) for i in range(11)}}]}
    b, c = _computed(t)
    assert len(findings.chart_bars(b, c)) == 12
    assert findings.chart_png(b, c)
