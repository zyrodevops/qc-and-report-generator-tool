"""
The defect table prints only the columns that hold data, and never the old
joined-percentage column that stretched every row.
"""

import io

import docx
from bs4 import BeautifulSoup

from app.compute.arithmetic import compute_table
from app.render.docx.engine import render_table
from app.render.html.engine import render_table_html
from app.render.table_columns import visible_columns

BLOCK = {
    "type": "table",
    "title": "THE CONDITION FOUND OF APPLE FRUITS:",
    "unit": "pcs",
    "grouping_label": "Count / Size",
    "categories": [
        {"key": "sound", "label": "Sound"},
        {"key": "pressure", "label": "Pressure"},   # nothing entered anywhere
        {"key": "rotten", "label": "Rotten"},       # an explicit zero
        {"key": "russet", "label": "Russet"},
    ],
    "rows": [
        {"group": "120", "values": {"sound": "67", "rotten": "0", "russet": "16"}},
        {"group": "120", "values": {"sound": "81", "rotten": "1", "russet": "11"}},
    ],
}


def test_empty_column_is_left_out_but_zero_is_kept():
    keys = [c["key"] for _, c in visible_columns(BLOCK)]
    assert keys == ["sound", "rotten", "russet"]


def test_table_with_no_rows_keeps_every_heading():
    empty = {**BLOCK, "rows": []}
    assert len(visible_columns(empty)) == 4


def test_html_and_docx_print_the_same_columns_and_no_joined_percentages():
    computed = compute_table(BLOCK)

    soup = BeautifulSoup(render_table_html(BLOCK, computed), "html.parser")
    html_header = [c.get_text(strip=True) for c in soup.find("tr").find_all(["th", "td"])]

    d = docx.Document()
    render_table(d, BLOCK, computed)
    buf = io.BytesIO()
    d.save(buf)
    table = docx.Document(io.BytesIO(buf.getvalue())).tables[0]
    docx_header = [c.text.strip() for c in table.rows[0].cells]

    assert html_header == docx_header == ["Count / Size", "Sound", "Rotten", "Russet", "Total (pcs)"]

    for tr in soup.find_all("tr"):
        for cell in tr.find_all(["td", "th"]):
            assert " / " not in cell.get_text() or cell.get_text(strip=True) == "Count / Size"
