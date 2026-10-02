import io
from docx import Document
from app.seeds.defaults import get_default_block_state
from app.render.docx.engine import render_docx
from app.ingest.documents.merge import particulars


def test_particulars_single_container_layout():
    state = get_default_block_state("fruit", "APPLE")
    docx_bytes = render_docx(state)
    assert len(docx_bytes) > 0

    doc = Document(io.BytesIO(docx_bytes))
    assert len(doc.tables) > 0

    t1 = doc.tables[0]
    assert len(t1.columns) == 5

    labels = [r.cells[0].text.strip() for r in t1.rows if r.cells[0].text.strip()]
    expected_fields = [
        "Policy No.",
        "Insurer",
        "Sum Insured",
        "Shipper",
        "Consignees",
        "Comm. Invoice No.",
        "Invoice Value",
        "Bill of Lading No.",
        "Vessel Name",
        "Voyage as per B/L",
        "Date of arrival",
        "Nature of Packing",
    ]
    for field in expected_fields:
        assert field in labels, f"Field '{field}' missing from Table 1"

    assert any("Consignment" in l for l in labels)
    assert any("Container" in l for l in labels)


def test_particulars_multi_container_layout():
    ship_multi = {
        "policy_number": "ANAJA9924226EGD29419 / BNAJA9924226EGD43804",
        "insurer": "China Pacific Property Insurance Co., Ltd.",
        "insured_value": "USD 36,036.00",
        "shipper": {"shown": "Dalian Micro Cold Agriculture Products Co., Ltd."},
        "consignee": {"shown": "Firangi Fresh\nB-153, New Fruit Market, Azadpur, Delhi"},
        "invoice_number": "WLN-PLUM-2026-001",
        "invoice_date": "21 July 2026",
        "invoice_value": "CFR Nhava Sheva USD 32,760.00",
        "document_number": "OOLU2334388420",
        "document_date": "8 August 2026",
        "vessel": "OOCL MALAYSIA",
        "voyage": "072W",
        "port_of_loading": "Dalian, China",
        "port_of_discharge": "Nhava Sheva, India",
        "arrival_date": "09 September 2026 at 10:23 IST",
        "containers": [
            {"container": "FBIU5161393", "type": "40RH", "seal": "OOLHP80993", "gross_kg": "28,350", "from": "B/L"},
            {"container": "OTPU6186064", "type": "40RH", "seal": "OOLHP80994", "gross_kg": "28,350", "from": "B/L"},
        ],
        "consignment": [
            {"container": "FBIU5161393", "variety": "Fresh Plum", "count": "50", "cartons": 320},
            {"container": "FBIU5161393", "variety": "Fresh Plum", "count": "55", "cartons": 2252},
            {"container": "OTPU6186064", "variety": "Fresh Plum", "count": "50", "cartons": 395},
            {"container": "OTPU6186064", "variety": "Fresh Plum", "count": "55", "cartons": 2234},
        ],
        "total_packages": 5201,
        "total_net_kg": "50,400",
        "total_gross_kg": "56,700",
        "sources": {},
    }
    fruit_particulars = particulars(ship_multi, general_cargo=False)
    assert len(fruit_particulars) == 15

    multi_state = get_default_block_state("fruit", "APPLE")
    for b in multi_state["blocks"]:
        if b["type"] == "particulars":
            b["rows"] = fruit_particulars

    multi_docx = render_docx(multi_state)
    multi_doc = Document(io.BytesIO(multi_docx))
    m_t1 = multi_doc.tables[0]
    m_labels = [r.cells[0].text.strip() for r in m_t1.rows if r.cells[0].text.strip()]

    assert any("FBIU5161393" in ml for ml in m_labels)
    assert any("OTPU6186064" in ml for ml in m_labels)


def test_particulars_headers_spelling_and_ordering():
    ship_single = {
        "containers": [{"container": "CONT001", "type": "40RH"}],
        "consignment": [{"container": "CONT001", "fruit": "Apple", "variety": "Royal Gala", "count": "100", "cartons": 500}],
    }
    p_single = particulars(ship_single, general_cargo=False)
    cons_row = next(r for r in p_single if r["label"] == "Consignment")
    assert cons_row["headers"] == ["Fresh Apple Variety", "Count / Size", "Total Boxes"]
    c_idx = p_single.index(cons_row)
    nop_idx = next(i for i, r in enumerate(p_single) if r["label"] == "Nature of Packing")
    assert nop_idx == c_idx + 1, "Nature of Packing must follow directly after Consignment"

    ship_multi = {
        "containers": [
            {"container": "CONT001", "type": "40RH"},
            {"container": "CONT002", "type": "40RH"},
        ],
        "consignment": [
            {"container": "CONT001", "variety": "Gala", "count": "100", "cartons": 500},
            {"container": "CONT002", "variety": "Fuji", "count": "120", "cartons": 600},
        ],
    }
    p_multi = particulars(ship_multi, general_cargo=False)
    labels = [r["label"] for r in p_multi]
    c1_idx = labels.index("Consignment # CONT001")
    c2_idx = labels.index("Consignment # CONT002")
    nop_idx = labels.index("Nature of Packing")
    assert c2_idx == c1_idx + 1
    assert nop_idx == c2_idx + 1
    assert p_multi[c1_idx]["headers"] == ["Commodity", "Sizes", "Total Boxes"]
    assert p_multi[c2_idx]["headers"] == ["Commodity", "Sizes", "Total Boxes"]
