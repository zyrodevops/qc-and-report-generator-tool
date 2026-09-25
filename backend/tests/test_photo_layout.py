"""
Survey photographs laid out the way the client's own photo tool lays them.

Measured in his reports: A4, two photos across, each 8.2 x 5.6 cm, no border,
captioned "Survey Photo No. N", 8 to a page (sometimes 6). The tool's options —
border and colour, caption wording, numbering start, caption font, size and
colour, compression — are kept on the block. Photos go in the right way up.
"""

import io
from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn
from httpx import ASGITransport, AsyncClient
from PIL import Image

from app.ingest.photos import generate_derived_copies, is_portrait, report_image
from app.main import app
from app.render import photo_layout as pl
from app.render.docx.engine import render_docx
from app.render.html.engine import render_photo_plate_html


def _jpeg(w, h, exif_orientation=None) -> bytes:
    img = Image.new("RGB", (w, h), (30, 90, 160))
    # Yellow marks the top-left corner of the picture as taken.
    img.paste((255, 255, 0), (0, 0, w // 3, h // 6))
    buf = io.BytesIO()
    if exif_orientation:
        exif = Image.Exif()
        exif[0x0112] = exif_orientation
        img.save(buf, format="JPEG", exif=exif)
    else:
        img.save(buf, format="JPEG")
    return buf.getvalue()


def _yellow_corner(path) -> str:
    im = Image.open(path).convert("RGB")
    w, h = im.size
    for name, xy in {"TL": (5, 5), "TR": (w - 5, 5), "BL": (5, h - 5), "BR": (w - 5, h - 5)}.items():
        r, g, b = im.getpixel(xy)
        if r > 200 and g > 200 and b < 80:
            return name
    return "?"


def test_defaults_are_the_clients_layout():
    lay = pl.layout_of({})
    assert lay["per_page"] == 8 and lay["border"] is False
    assert lay["caption_keyword"] == "Survey Photo No." and lay["caption_font"] == "Arial"
    assert lay["caption_size"] == 11 and lay["number_from"] == 1
    assert round(pl.IMAGE_W_PX * pl.CM_PER_PX, 2) == 8.2 and round(pl.IMAGE_H_PX * pl.CM_PER_PX, 2) == 5.56


def test_bad_settings_fall_back_to_the_default():
    lay = pl.layout_of({"layout": {"per_page": 5, "caption_size": 99, "caption_color": "#123456",
                                   "quality": "huge", "number_from": 0, "border": "yes"}})
    assert (lay["per_page"], lay["caption_size"], lay["caption_color"], lay["quality"],
            lay["number_from"], lay["border"]) == (8, 11, "000000", "balanced", 1, False)
    ok = pl.layout_of({"layout": {"per_page": "6", "caption_color": "#000080", "number_from": "21"}})
    assert (ok["per_page"], ok["caption_color"], ok["number_from"]) == (6, "000080", 21)


def test_typed_numbers_are_not_printed_twice():
    assert pl.strip_number("Photo 7") == ""
    assert pl.strip_number("Survey Photo No. 3: Seal intact") == "Seal intact"
    assert pl.strip_number("Photo No. 5 — Cargo stowage") == "Cargo stowage"
    assert pl.strip_number("Container exterior") == "Container exterior"
    lay = pl.layout_of({})
    assert pl.caption(lay, 7, "Photo 7") == "Survey Photo No. 7"
    assert pl.caption(lay, 2, "Seal intact") == "Survey Photo No. 2 — Seal intact"


def test_pages_hold_the_chosen_number_and_a_heading_takes_a_row():
    photos = list(range(19))
    assert [len(p) for p in pl.pages(photos, pl.layout_of({}))] == [8, 8, 3]
    six = pl.layout_of({"layout": {"per_page": 6, "show_heading": True}})
    assert [len(p) for p in pl.pages(photos, six)] == [4, 6, 6, 3]


def test_photos_come_out_the_right_way_up(tmp_path):
    # Stored sideways with "rotate 90" in the EXIF, as a phone saves it.
    sideways = Image.open(io.BytesIO(_jpeg(400, 300))).rotate(90, expand=True)
    buf = io.BytesIO()
    exif = Image.Exif()
    exif[0x0112] = 6
    sideways.save(buf, format="JPEG", exif=exif)
    raw = buf.getvalue()
    assert not is_portrait(raw)  # upright it is landscape
    d = generate_derived_copies(raw, "x", tmp_path)
    assert Image.open(d["report"]).size == (400, 300) and _yellow_corner(d["report"]) == "TL"

    tall = _jpeg(300, 400)
    assert is_portrait(tall)
    turned = generate_derived_copies(tall, "y", tmp_path, rotation=90)
    assert Image.open(turned["report"]).size == (400, 300)
    assert _yellow_corner(turned["report"]) == "BL"  # a quarter turn anticlockwise


def test_report_copy_is_made_from_the_original_and_kept(tmp_path):
    src = tmp_path / "p_original.jpg"
    src.write_bytes(_jpeg(3000, 2000))
    before = src.read_bytes()
    asset = {"original_path": str(src), "rotation": 0}
    out = report_image(asset, tmp_path / "d", 1600, 65, "small")
    assert Image.open(out).size == (1600, 1067)
    assert report_image(asset, tmp_path / "d", 1600, 65, "small") == out
    assert src.read_bytes() == before  # never touched


def _state(tmp_path, n, layout=None, after=True):
    assets, groups = {}, []
    for i in range(1, n + 1):
        p = tmp_path / f"s{i}_original.jpg"
        p.write_bytes(_jpeg(800, 600))
        assets[f"a{i}"] = {"id": f"a{i}", "original_path": str(p), "derived_paths": {}, "rotation": 0}
        groups.append({"id": f"g{i}", "observation": "", "asset_ids": [f"a{i}"]})
    blocks = [
        {"id": "n1", "type": "narrative", "section": "OUR SURVEY", "additional_text": "Before."},
        {"id": "ph", "type": "photo_plate", "series_id": "survey", "label": "SURVEY PHOTOGRAPHS:",
         "groups": groups, "layout": layout or {}},
    ]
    if after:
        blocks.append({"id": "n2", "type": "narrative", "section": "CAUSE", "additional_text": "After."})
    return {"metadata": {"family": "QC_REPORT"}, "assets": assets, "blocks": blocks}


def test_word_layout(tmp_path, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "DERIVED_DIR", tmp_path / "derived")

    d = docx.Document(io.BytesIO(render_docx(_state(tmp_path, 19))))
    photo_sec = d.sections[1]
    assert round(photo_sec.page_width.cm, 1) == 21.0 and round(photo_sec.page_height.cm, 1) == 29.7
    assert round(photo_sec.left_margin.cm, 2) == 2.54 and round(photo_sec.top_margin.cm, 2) == 2.54
    # Text after the photos goes back to the report's own margins.
    assert d.sections[-1].left_margin == d.sections[0].left_margin
    photo_tables = [t for t in d.tables if t._tbl.xpath(".//pic:pic")]
    assert [len(t._tbl.xpath(".//pic:pic")) for t in photo_tables] == [8, 8, 3]
    shape = d.inline_shapes[0]
    assert (shape.width, shape.height) == (310 * 9525, 210 * 9525)
    caps = [c.text for t in photo_tables for r in t.rows for c in r.cells if c.text]
    assert caps[:3] == ["Survey Photo No. 1", "Survey Photo No. 2", "Survey Photo No. 3"]
    run = photo_tables[0].rows[1].cells[0].paragraphs[0].runs[0]
    assert run.font.name == "Arial" and run.font.size.pt == 11
    assert not photo_tables[0]._tbl.xpath(".//w:tcBorders/w:top[@w:val='single']")
    assert photo_tables[0].rows[0]._tr.trPr.find(qn("w:cantSplit")) is not None

    # A report that ends with the photos gets no empty page after them.
    tail = docx.Document(io.BytesIO(render_docx(_state(tmp_path, 3, after=False))))
    assert len(tail.sections) == 2


def test_word_options(tmp_path, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "DERIVED_DIR", tmp_path / "derived")
    layout = {"per_page": 6, "border": True, "border_color": "000080", "caption_keyword": "Cargo Photo No.",
              "number_from": 21, "caption_font": "Times New Roman", "caption_size": 12, "caption_color": "000080"}
    d = docx.Document(io.BytesIO(render_docx(_state(tmp_path, 7, layout))))
    photo_tables = [t for t in d.tables if t._tbl.xpath(".//pic:pic")]
    assert [len(t._tbl.xpath(".//pic:pic")) for t in photo_tables] == [6, 1]
    cell = photo_tables[0].rows[1].cells[0]
    run = cell.paragraphs[0].runs[0]
    assert cell.text == "Cargo Photo No. 21"
    assert (run.font.name, run.font.size.pt, str(run.font.color.rgb)) == ("Times New Roman", 12, "000080")
    assert cell._tc.xpath("./w:tcPr/w:tcBorders/w:top[@w:val='single'][@w:color='000080']")


def test_html_matches_word(tmp_path, monkeypatch):
    from app.compute.arithmetic import compute
    from app.config import settings
    monkeypatch.setattr(settings, "DERIVED_DIR", tmp_path / "derived")
    state = compute(_state(tmp_path, 9, {"border": True}))
    block = next(b for b in state["blocks"] if b["type"] == "photo_plate")
    out = render_photo_plate_html(block, block.get("_computed", {}), state["assets"])
    assert out.count('<table class="photo-table"') == 2 and out.count("page-break-before") == 1
    assert out.count("data:image/jpeg;base64") == 9
    assert "Survey Photo No. 9" in out and "1.5pt solid #000000" in out


@pytest.mark.asyncio
async def test_upload_turns_upright_photos_and_rotate_turns_them_back():
    transport = ASGITransport(app=app)
    made = []
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        tok = (await ac.post("/api/auth/login", json={"password": "surveyor123"})).json()["token"]
        hdr = {"Authorization": f"Bearer {tok}"}
        rep = await ac.post("/api/reports", json={"family": "QC_REPORT", "commodity": "APPLE",
                                                  "template_id": "perishable_qc_sea"}, headers=hdr)
        rid = rep.json()["id"]
        try:
            files = [("files", ("tall.jpg", _jpeg(300, 400), "image/jpeg")),
                     ("files", ("wide.jpg", _jpeg(400, 300), "image/jpeg"))]
            res = await ac.post(f"/api/reports/{rid}/assets/photos/batch", files=files,
                                data={"landscape": "true"}, headers=hdr)
            assert res.status_code == 201, res.text
            tall, wide = res.json()["assets"]
            made += [tall, wide]
            assert (tall["rotation"], wide["rotation"]) == (90, 0)
            assert tall["url"].endswith("?v=r90")
            assert Image.open(tall["derived_paths"]["report"]).size == (400, 300)
            original = Path(tall["original_path"]).read_bytes()

            res = await ac.post(f"/api/reports/{rid}/assets/{tall['id']}/rotate", json={"turns": -1}, headers=hdr)
            assert res.status_code == 200, res.text
            entry = res.json()["asset"]
            assert entry["rotation"] == 0 and not entry["url"].endswith("?v=r90")
            assert Image.open(entry["derived_paths"]["report"]).size == (300, 400)
            assert Path(tall["original_path"]).read_bytes() == original  # original untouched
            state = (await ac.get(f"/api/reports/{rid}", headers=hdr)).json()["block_state"]
            assert state["assets"][tall["id"]]["rotation"] == 0

            bad = await ac.post(f"/api/reports/{rid}/assets/{tall['id']}/rotate", json={"turns": 2}, headers=hdr)
            assert bad.status_code == 400

            # Without the tick, an upright photo stays upright.
            res = await ac.post(f"/api/reports/{rid}/assets/photos/batch",
                                files=[("files", ("t2.jpg", _jpeg(300, 400), "image/jpeg"))], headers=hdr)
            made += res.json()["assets"]
            assert res.json()["assets"][0]["rotation"] == 0
        finally:
            await ac.delete(f"/api/reports/{rid}", headers=hdr)
            for a in made:
                for f in [a["original_path"], *a["derived_paths"].values()]:
                    Path(f).unlink(missing_ok=True)
