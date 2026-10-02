"""
General cargo tables: weighbridge, tally, WEIGHT FINAL SUMMARY, containers &
seals, weather, summary of reserve. All figures made up for the test.
"""

import io

import docx

from app.compute.arithmetic import compute
from app.compute.gc_tables import fmt, num, report_table, tally, unit_segments, weighbridge, weight_summary
from app.compute.traceability import check_traceability
from app.render.docx.engine import render_docx
from app.render.html.engine import render_html
from app.seeds import general_cargo as gcmod
from app.seeds.defaults import get_default_block_state


def _state():
    return get_default_block_state("general_cargo_sea_survey", "BAGGED_FOOD")


def _by_id(s):
    return {b["id"]: b for b in s["blocks"]}


W1 = {"weighbridge": "Test Weighbridge", "slip_no": "4521", "date": "5 June 2026", "gross": "34,560",
      "tare_truck": "8,200", "tare_container": "3,800", "declared": "23,000", "basis": "B/L"}


def test_numbers_are_read_and_printed_as_entered():
    assert num("34,560") == 34560 and num("") is None and num("abc") is None
    assert fmt(num("22560")) == "22,560" and fmt(num("12.50")) == "12.5" and fmt(num("-440")) == "-440"


def test_weighbridge_works_out_the_found_weight_and_the_shortage():
    w = weighbridge({"weights": W1})
    assert w["found"] == 22560 and w["diff"] == -440
    assert w["rows"] == [
        ["Weighbridge", "Test Weighbridge"],
        ["Weight Slip No. and date", "4521 dated 5 June 2026"],
        ["Gross weight found of container + cargo + trailer", "34,560"],
        ["Less: Tare weight of trailer", "8,200"],
        ["Less: Marked tare weight of container", "3,800"],
        ["Found gross weight of cargo", "22,560"],
        ["Less: Gross weight of cargo as per Bill of Lading", "23,000"],
        ["Difference (shortage)", "440"],
    ]
    # nothing weighed: no table
    assert weighbridge({}) is None and weighbridge({"weights": {"gross": ""}}) is None


def test_tally_and_its_total():
    t = tally({"tally": {"basis": "Packing List", "rows": [
        {"item": "Bags", "document": "500", "sound": "440", "damaged": "45"},
        {"item": "Cartons", "document": "20", "sound": "20", "damaged": ""}]}})
    assert t["columns"][1] == "As per Packing List"
    assert t["rows"][0] == ["Bags", "500", "440", "45", "-15"]
    assert t["rows"][1] == ["Cartons", "20", "20", "", "0"]
    assert t["rows"][-1] == ["Total", "520", "460", "45", "-15"]


def test_weight_final_summary_needs_two_containers():
    s = _state()
    u1 = _by_id(s)["b_survey_1"]
    u1.update(container="TSTU1234565", weights=W1)
    assert weight_summary(s["blocks"]) is None
    u2 = gcmod.survey_unit(2, container="TSTU7654320")
    u2["weights"] = {"gross": "21,000", "declared": "21,050"}
    s["blocks"].insert(s["blocks"].index(u1) + 1, u2)
    ws = weight_summary(s["blocks"])
    assert ws["title"] == "WEIGHT FINAL SUMMARY"
    assert ws["columns"][1:3] == ["Ascertained weight (kg)", "Weight as per B/L (kg)"]
    assert ws["rows"] == [["TSTU1234565", "22,560", "23,000", "-440"],
                          ["TSTU7654320", "21,000", "21,050", "-50"],
                          ["Total", "43,560", "44,050", "-490"]]


def test_report_tables():
    seals = report_table({"kind": "seals", "rows": [
        {"container": "TSTU1234565", "size": "40' HC", "seal_doc": "AB 123456", "seal_found": "ab123456"},
        {"container": "TSTU7654320", "size": "20'", "seal_doc": "CD 1", "seal_found": "CD 2"},
        {"container": "", "size": "", "seal_doc": "", "seal_found": ""}]})
    assert [r[-1] for r in seals["rows"]] == ["Yes", "No"] and seals["title"] == ""
    reserve = report_table({"kind": "reserve", "currency": "USD", "rows": [
        {"description": "45 bags wet", "quantity": "2,250 kg", "value": "1,125.50"},
        {"description": "15 bags torn", "quantity": "750 kg", "value": "375"}]})
    assert reserve["title"] == "SUMMARY OF RESERVE" and reserve["columns"][2] == "Invoice value (USD)"
    assert reserve["rows"][-1] == ["Total", "", "1,500.5"]
    assert report_table({"kind": "weather", "rows": []}) is None


def test_tables_go_where_marked_or_after_the_text():
    segs = unit_segments("Lead in.\n\n(weight table)\n\nAfter the weights.")
    assert segs == [("text", "Lead in."), ("table", "weights"), ("text", "After the weights."),
                    ("table", "damage"), ("table", "tally")]


def test_word_and_html_carry_the_tables_and_every_figure_is_traceable():
    s = _state()
    b = _by_id(s)
    u1 = b["b_survey_1"]
    u1.update(container="TSTU1234565", weights=W1,
              additional_text="The weight details noted are as follows:\n\n(weight table)",
              tally={"rows": [{"item": "Bags", "document": "500", "sound": "440", "damaged": "45"}]})
    u2 = gcmod.survey_unit(2, container="TSTU7654320")
    u2["weights"] = {"gross": "21,000", "declared": "21,050"}
    s["blocks"].insert(s["blocks"].index(u1) + 1, u2)
    b["b_reserve"].update(included=True, currency="INR",
                          rows=[{"description": "Wet bags", "quantity": "45", "value": "90,000"}])
    b["b_seals"].update(included=True, rows=[{"container": "TSTU1234565", "seal_doc": "AB1", "seal_found": "AB1"}])

    data = render_docx(s)
    d = docx.Document(io.BytesIO(data))
    paras = [p.text for p in d.paragraphs]
    heads = [t.rows[0].cells[0].text for t in d.tables]
    assert "Particulars" in heads and "Description" in heads and "Container No." in heads
    assert "WEIGHT FINAL SUMMARY" in paras and "SUMMARY OF RESERVE" in paras
    assert "(weight table)" not in paras
    weights = next(t for t in d.tables if t.rows[0].cells[0].text == "Particulars")
    assert weights.rows[-1].cells[1].text == "440"
    # the check on printed numbers finds every worked-out figure
    result = check_traceability(compute(s), data)
    assert result.passed, result.untraceable

    h = render_html(s)
    assert "weights-table" in h and "weight-summary" in h and "gc-reserve-table" in h and "90,000" in h
    # a table left out does not print
    b["b_reserve"]["included"] = False
    assert "gc-reserve-table" not in render_html(s)


def test_new_reports_have_the_tables_switched_off():
    s = _state()
    tables = [x for x in s["blocks"] if x["type"] == "gc_table"]
    assert [t["kind"] for t in tables] == ["seals", "weather", "reserve"]
    assert all(t["included"] is False for t in tables)
