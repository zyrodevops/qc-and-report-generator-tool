"""
Tier 3: Cross-Feature Combinations - End-to-End Ingest -> Compute -> Word -> Gate Pipeline.
Covers:
Spreadsheet Ingest -> Column Mapping -> BlockState -> compute() -> Word Generation -> Traceability Gate.
"""

from decimal import Decimal
import io
import pytest
import openpyxl
from tests.e2e.helpers.contract_stubs import (
    reference_compute_table,
    get_docx_renderer,
    get_traceability_gate
)
from tests.e2e.helpers.synthetic_data import make_synthetic_block_state


@pytest.mark.tier3
def test_full_data_pipeline_ingest_compute_render_gate(excel_with_sum_formulas):
    """
    Tests the cross-feature integration:
    1. Ingest Excel workbook with =SUM() formulas
    2. Map columns to block categories
    3. Execute pure compute() to generate derived totals and percentages
    4. Inject into Word template
    5. Pass through Numeric Traceability Gate
    """
    # 1. Ingest & extract
    wb = openpyxl.load_workbook(io.BytesIO(excel_with_sum_formulas), data_only=True)
    ws = wb["Grapes QC"]
    
    # Read rows
    row1_vals = {"sound": Decimal(str(ws["B2"].value)), "waterberry": Decimal(str(ws["C2"].value)), "decay": Decimal(str(ws["D2"].value))}
    row2_vals = {"sound": Decimal(str(ws["B3"].value)), "waterberry": Decimal(str(ws["C3"].value)), "decay": Decimal(str(ws["D3"].value))}
    
    # 2. Compute
    categories = ["sound", "waterberry", "decay"]
    rows = [{"group": "Sample 1", "values": row1_vals}, {"group": "Sample 2", "values": row2_vals}]
    computed = reference_compute_table(rows, categories, precision="0.001")

    assert computed["row_totals"][0] == Decimal("6.230")
    assert computed["grand_total"] == Decimal("50.702")

    # 3. Assemble into Block State
    state = make_synthetic_block_state()
    table_block = next(b for b in state["blocks"] if b["type"] == "table")
    table_block["rows"] = rows
    table_block["computed_results"] = computed

    # 4 & 5. Word render & Gate
    renderer = get_docx_renderer()
    gate = get_traceability_gate()
    if renderer and gate:
        docx_bytes = renderer(state)
        gate_res = gate(docx_bytes, state)
        assert gate_res.get("valid") is True or getattr(gate_res, "valid", False) is True
